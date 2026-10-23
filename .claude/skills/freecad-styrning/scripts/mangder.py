"""Mängdunderlag ur modellen: mängder per del och totalt, som underlag för tillverkarens kalkyl. Inga priser:
prissättningen görs av tillverkaren med ett eget verktyg, utanför skillen.

  python3 mangder.py <modell.FCStd> [--antal 1] [--ut <mapp>]

Mängderna räknas ur geometrin (FreeCAD), för en sammanställning (delar med positionsnummer) eller en enstaka
del (Part Design, SheetMetal):
  - vikt per material (densitet efter materialet: stål, rostfritt, aluminium, mässing, glas, trä)
  - yta (m2) för ytbehandling: lack, förzinkning, patinering
  - kap: raka kap (två per stång/profil) och geringar (sneda kap, Pocket "Gering")
  - hål per diameter (hålgrupperna), borrade i profiler och skurna i plåt
  - laser: skärlängd (m) och antal håltagningar (insticksställen) för plåtar (PL med bredd >= 100) och
    bockade delar (utbredningens kontur)
  - bockningar i bockade delar (SheetMetal)
  - svets: uppskattad längd (m) = kontaktkanterna mellan delar som rör varandra (yttre konturen av varje
    kontaktyta). Uppskattning: verkligt svetsomfång bestäms av konstruktören.
  - inköpta detaljer (k.inkop) med antal
Ger `<modell>_mangder.csv` (öppnas i Excel) och `<modell>_mangder.json` (till beställningspaketet, paket.py).

Kalkylunderlag för ett projekt med många sammanställningar:

  python3 mangder.py samla <modell.FCStd eller _mangder.json>[:antal] ... --namn <projekt> [--ut <mapp>]

samlar mängderna (antalet ur `:antal`, annars ur json-filens antal) i `<projekt> kalkylunderlag.xlsx` med bladen
Mängdspec (per märke: antal, vikt, yta, svets, hål, laser, kap, delar), Kaplista profiler (färdiga längder med
antal och löpmeter, summa per profil), Plåtlista (tjocklek x bredd x längd, summa per tjocklek) och Per material.
Saknas openpyxl blir det tre CSV-filer i stället. Saknas en modells _mangder.json körs mängdberäkningen först.
Underlaget är maskinellt uttaget ur modellerna och ska granskas: bladen har en rad om det.
"""
import os, sys, re, json, argparse, collections

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---------------------------------------------------------------- inne i FreeCAD
def _platta(s):
    """Skuren platta: prismatisk kropp (plana ytor parallella eller vinkelräta mot den största ytan, cylindrar
    vinkelräta) med tjockleken under halva minsta sidan."""
    plan = [f for f in s.Faces if f.Surface.TypeId == "Part::GeomPlane"]
    if not plan:
        return False
    n = max(plan, key=lambda f: f.Area).normalAt(0, 0)
    for f in s.Faces:
        if f.Surface.TypeId == "Part::GeomPlane":
            k = abs(f.normalAt(0, 0).dot(n))
            if 1e-6 < k < 1 - 1e-6:
                return False
        elif f.Surface.TypeId == "Part::GeomCylinder":
            if abs(abs(f.Surface.Axis.dot(n)) - 1) > 1e-6:
                return False
        else:
            return False
    proj = [v.Point.dot(n) for v in s.Vertexes]
    t = max(proj) - min(proj)
    bb = s.BoundBox
    sidor = sorted((bb.XLength, bb.YLength, bb.ZLength))[1:]
    return t < 0.5 * min(sidor)


def mangder(doc, logg=print):
    import fcapi
    import Part
    import stal
    g = doc.getObject("Sammanstallning")
    if g is not None:
        ds = [o for o in g.Group if hasattr(o, "Pos")]
    else:                                        # enstaka del: slutliga solider (Body, SheetMetal), inte utbredningen
        kand = [o for o in doc.Objects if o.Name != "Utbredning" and getattr(o, "Shape", None) is not None and
                not o.TypeId.startswith(("TechDraw::", "Sketcher::", "Spreadsheet::", "App::", "PartDesign::Feature", "PartDesign::Pad",
                                         "PartDesign::Pocket", "PartDesign::Hole", "PartDesign::Fillet", "PartDesign::Chamfer",
                                         "PartDesign::Revolution", "PartDesign::Additive", "PartDesign::Subtractive", "PartDesign::Linear",
                                         "PartDesign::Polar", "PartDesign::Mirrored", "PartDesign::MultiTransform", "PartDesign::Groove",
                                         "PartDesign::Thickness", "PartDesign::Draft", "PartDesign::Boolean"))
                and o.Shape.Solids and (o.TypeId == "PartDesign::Body" or o.getParentGeoFeatureGroup() is None)]
        ds = [o for o in kand if not any(p in kand for p in o.InList)]
    per_del, vikt_mat = [], {}
    tot = dict(kap=0, geringar=0, hal=0, laser_m=0.0, instick=0, bockningar=0, yta_m2=0.0, vikt=0.0, delar=len(ds))
    halper, inkop = {}, []
    for o in ds:
        if hasattr(o, "Vikt") and getattr(o, "Material", "") == "INKÖP":     # inköpt detalj (k.inkop)
            inkop.append(dict(pos=o.Pos, text=o.Profil, vikt=float(o.Vikt or 0), kalla=getattr(o, "Kalla", "")))
            tot["delar"] -= 1
            continue
        mat = getattr(o, "Material", "") or "Stål"
        profil = getattr(o, "Profil", "") or o.Label
        s = o.Shape
        vikt = float(getattr(o, "Vikt", 0) or 0) or s.Volume * 1e-9 * stal.densitet(mat)
        info = None
        try:
            info = stal.tolka(profil, kallor=False)
        except Exception:
            pass
        feat = getattr(o, "Group", [])
        geringar = sum(1 for f in feat if f.TypeId == "PartDesign::Pocket" and f.Name.startswith("Gering"))
        hal = 0
        for f in feat:
            if f.TypeId == "PartDesign::Pocket" and hasattr(f, "Halgrupp"):
                sk = f.Profile[0] if isinstance(f.Profile, (list, tuple)) else f.Profile
                for geo in sk.Geometry:
                    if isinstance(geo, Part.Circle):
                        hal += 1
                        d = round(geo.Radius * 2, 1)
                        halper[d] = halper.get(d, 0) + 1
        bockad = o.TypeId != "PartDesign::Body"
        plat = bockad or (info is not None and info["typ"] == "PL" and profil.upper().startswith("PL") and info["b"] >= 100) or (
            not hasattr(o, "Profil") and _platta(s))
        laser, instick, bock = 0.0, 0, 0
        if getattr(o, "Bockradie", 0):                   # bockat rör (k.rorbage): bågen bockas i en operation
            bock = 1
        if plat:
            plan = [f for f in s.Faces if f.Surface.TypeId == "Part::GeomPlane"]
            f = max(plan, key=lambda f: f.Area) if plan else None
            if bockad:                                   # bockad plåt: utbredningens kontur och bockarna
                try:
                    import plat as P
                    utb = P.utbred(s, P.rotyta(s, f.normalAt(0, 0)))
                    f, bock = utb["face2d"], len(utb["bockar2d"])
                except Exception as e:
                    logg("OBS: %s kunde inte bredas ut för mängdunderlaget (%s)" % (o.Label, e))
                    bock = sum(1 for f_ in s.Faces if f_.Surface.TypeId == "Part::GeomCylinder") // 2
            if f is not None:
                laser = sum(w.Length for w in f.Wires) / 1000.0
                instick = len(f.Wires)
        kap = 0 if plat else (2 if info is not None and info["typ"] != "PL" or profil.upper().startswith(("FL", "FS", "PL")) else 0)
        kap = max(kap - min(geringar, 2), 0)
        yta = s.Area / 1e6
        per_del.append(dict(namn=fcapi.etikett(o), pos=getattr(o, "Pos", ""), profil=profil, material=mat, vikt=vikt, yta_m2=yta,
                            langd=stal.langd(o) if o.TypeId == "PartDesign::Body" else 0.0, kap=kap, geringar=geringar,
                            hal=hal, laser_m=laser, instick=instick, bockningar=bock))
        vikt_mat[mat] = vikt_mat.get(mat, 0.0) + vikt
        for k, v in (("kap", kap), ("geringar", geringar), ("hal", hal), ("laser_m", laser), ("instick", instick),
                     ("bockningar", bock), ("yta_m2", yta), ("vikt", vikt)):
            tot[k] += v
    # svets: kontaktkanter mellan delar som rör varandra (yttre slingan av varje kontaktyta)
    svets, fogar = 0.0, 0
    for i, o1 in enumerate(ds):
        b1 = o1.Shape.BoundBox
        b1.enlarge(0.2)
        for o2 in ds[i + 1:]:
            if not b1.intersect(o2.Shape.BoundBox):
                continue
            try:
                if o1.Shape.distToShape(o2.Shape)[0] > 0.05:
                    continue
                kanter = o1.Shape.section(o2.Shape).Edges
            except Exception:
                continue
            if not kanter:
                continue
            slingor = [Part.Wire(w) for w in Part.__sortEdges__(kanter) if w] if len(kanter) > 1 else [Part.Wire(kanter)]
            yttre = []
            for w in slingor:                               # släpp slingor som ligger inuti en annan (rörets insida)
                bw = w.BoundBox
                if any(w2 is not w and w2.BoundBox.isInside(bw.getPoint(0)) and w2.BoundBox.isInside(bw.getPoint(6))
                       and w2.BoundBox.DiagonalLength > bw.DiagonalLength + 0.1 for w2 in slingor):
                    continue
                yttre.append(w)
            L = sum(w.Length for w in yttre)
            if L > 1.0:
                svets += L
                fogar += 1
    tot["svets_m"] = svets / 1000.0
    tot["fogar"] = fogar
    tot["inkopta"] = len(inkop)
    return dict(marke=getattr(g, "Marke", "") if g else "", benamning=getattr(g, "Namn", "") if g else "",
                per_del=per_del, inkop=inkop, totalt=tot, vikt_per_material=vikt_mat, hal_per_diameter={str(k): v for k, v in sorted(halper.items())},
                **{k: tot[k] for k in tot})


def _jobb(a):
    import FreeCAD as App
    from fcapi import logg
    doc = App.openDocument(a["fil"])
    doc.recompute()
    m = mangder(doc, logg)
    t = m["totalt"]
    logg("Delar %d (+ %d inköpta), vikt %.1f kg (%s)" % (t["delar"], t["inkopta"], t["vikt"], ", ".join("%s %.1f kg" % kv for kv in sorted(m["vikt_per_material"].items()))))
    logg("Yta %.2f m2, kap %d, geringar %d, hål %d (%s), laser %.2f m / %d håltagningar, bockningar %d, svets ca %.1f m i %d fogar" % (
        t["yta_m2"], t["kap"], t["geringar"], t["hal"], ", ".join("Ø%s: %d" % kv for kv in m["hal_per_diameter"].items()) or "-",
        t["laser_m"], t["instick"], t["bockningar"], t["svets_m"], t["fogar"]))
    return m


def _efterbehandla(m, a):
    utmapp = os.path.abspath(a["ut"]) if a.get("ut") else os.path.dirname(a["fil"])
    os.makedirs(utmapp, exist_ok=True)
    bas = os.path.join(utmapp, os.path.splitext(os.path.basename(a["fil"]))[0])
    antal = int(a.get("antal") or 1)
    m.update(antal=antal)
    for k in ("ok", "dialoger"):
        m.pop(k, None)
    with open(bas + "_mangder.json", "w", encoding="utf-8") as f:
        json.dump(m, f, ensure_ascii=False, indent=1)
    tal = lambda x, d=2: ("%.*f" % (d, x)).replace(".", ",")
    with open(bas + "_mangder.csv", "w", encoding="utf-8-sig") as f:
        f.write("MÄNGDUNDERLAG;%s;antal %d\n\n" % (os.path.basename(a["fil"]), antal))
        f.write("DEL;POS;PROFIL;MATERIAL;LÄNGD mm;VIKT kg;YTA m2;KAP;GERINGAR;HÅL;LASER m;HÅLTAGNINGAR;BOCKNINGAR\n")
        for d in m["per_del"]:
            f.write("%s;%s;%s;%s;%s;%s;%s;%d;%d;%d;%s;%d;%d\n" % (d["namn"], d["pos"], d["profil"], d["material"], tal(d["langd"], 0),
                                                              tal(d["vikt"]), tal(d["yta_m2"], 3), d["kap"], d["geringar"], d["hal"],
                                                              tal(d["laser_m"]), d["instick"], d["bockningar"]))
        t = m["totalt"]
        f.write("\nTOTALT PER ST;;;;;%s;%s;%d;%d;%d;%s;%d;%d\n" % (tal(t["vikt"]), tal(t["yta_m2"], 3), t["kap"], t["geringar"], t["hal"],
                                                            tal(t["laser_m"]), t["instick"], t["bockningar"]))
        f.write("SVETS (UPPSKATTAD);%s m;%d fogar\n" % (tal(t["svets_m"]), t["fogar"]))
        for mat, kg in sorted(m["vikt_per_material"].items()):
            f.write("VIKT %s;%s kg\n" % (mat, tal(kg)))
        for d, n in m["hal_per_diameter"].items():
            f.write("HÅL Ø%s;%d st\n" % (d, n))
        if m.get("inkop"):
            f.write("\nINKÖP;POS;ANTAL;KÄLLA\n")
            per = {}
            for x in m["inkop"]:
                k_ = (x["text"], x["pos"], x.get("kalla", ""))
                per[k_] = per.get(k_, 0) + 1
            for (text, pos, kalla), n in sorted(per.items()):
                f.write("%s;%s;%d;%s\n" % (text, pos, n * antal, kalla))
    print("Mängdunderlag: %s, %s" % (bas + "_mangder.csv", os.path.basename(bas + "_mangder.json")))
    return m


# ---------------------------------------------------------------- samla (vanlig Python)
GRANSKA = "Maskinellt uttaget ur modellerna. Granskas före kalkyl och beställning."


def samla(filer, namn, ut=None, logg=print):
    """Kalkylunderlag för flera sammanställningar: filer = ["SB1.FCStd:8", "SP1_mangder.json", ...]."""
    mangder_ = []
    for f in filer:
        antal = None
        m_ = re.match(r"^(.*?):(\d+)$", f)
        if m_:
            f, antal = m_.group(1), int(m_.group(2))
        f = os.path.abspath(f)
        js = f if f.endswith(".json") else os.path.splitext(f)[0] + "_mangder.json"
        if not os.path.exists(js):
            import fcapi
            a = dict(fil=f, antal=antal or 1, ut=None)
            _efterbehandla(fcapi.kor(os.path.abspath(__file__), a), a)
        d = json.load(open(js, encoding="utf-8"))
        marke = d.get("marke") or os.path.basename(js).replace("_mangder.json", "")
        mangder_.append((marke, antal if antal is not None else int(d.get("antal") or 1), d))
    rader, kap, plat, mat = [], collections.defaultdict(lambda: [0, 0.0, set()]), \
        collections.defaultdict(lambda: [0, 0.0, set()]), collections.Counter()
    for marke, n, d in mangder_:
        t = d["totalt"]
        rader.append([marke, d.get("benamning", ""), n, round(t["vikt"], 1), round(t["vikt"] * n, 1), round(t["yta_m2"] * n, 2),
                      round(t["svets_m"] * n, 1), t["hal"] * n, round(t["laser_m"] * n, 2), t["kap"] * n, t["delar"] * n])
        for k_, v in d["vikt_per_material"].items():
            mat[k_] += v * n
        for p in d["per_del"]:
            ref = "%s %s" % (marke, p["pos"])
            m2 = re.match(r"^PL([\d.]+)\*([\d.]+)$", p["profil"].upper())
            if m2 and not p.get("bockningar"):
                key = (float(m2.group(1)), float(m2.group(2)), round(p["langd"]), p["material"])
                tab = plat
            else:
                key = (p["profil"], p["material"], round(p["langd"]))
                tab = kap
            tab[key][0] += n
            tab[key][1] += p["vikt"] * n
            tab[key][2].add(ref)
    tot = [sum(r[i] for r in rader) for i in range(2, 11)]
    rader.append(["TOTALT", "", tot[0], "", round(tot[2], 1), round(tot[3], 2), round(tot[4], 1), tot[5], round(tot[6], 2), tot[7], tot[8]])
    rk = [[p, m, L, n, round(L * n / 1000, 2), round(v, 1), ", ".join(sorted(r))]
          for (p, m, L), (n, v, r) in sorted(kap.items(), key=lambda x: (x[0][0], -x[0][2]))]
    rp = [["PL%g*%g" % (t_, b), t_, b, L, m, n, round(v, 1), ", ".join(sorted(r))]
          for (t_, b, L, m), (n, v, r) in sorted(plat.items(), key=lambda x: (x[0][0], x[0][1], -x[0][2]))]
    perprof = collections.defaultdict(lambda: [0.0, 0.0])
    for r in rk:
        perprof[r[0]][0] += r[4]
        perprof[r[0]][1] += r[5]
    pertj = collections.defaultdict(lambda: [0.0, 0.0])
    for r in rp:
        pertj[r[1]][0] += r[2] * r[3] * r[5] / 1e6
        pertj[r[1]][1] += r[6]
    utmapp = os.path.abspath(ut or os.path.dirname(os.path.abspath(filer[0].split(":")[0])))
    os.makedirs(utmapp, exist_ok=True)
    blad = [
        ("Mängdspec", ["Märke", "Benämning", "Antal", "Vikt/st kg", "Vikt tot kg", "Yta tot m²", "Svets ca m", "Hål st",
                       "Laser m", "Kap st", "Delar st"], rader, []),
        ("Kaplista profiler", ["Profil", "Material", "Längd mm", "Antal", "Löpmeter", "Vikt kg", "Märke / pos"], rk,
         [["Summa per profil"]] + [[p, "", "", "", round(lm, 2), round(v, 1)] for p, (lm, v) in sorted(perprof.items())]),
        ("Plåtlista", ["Plåt", "t mm", "b mm", "L mm", "Material", "Antal", "Vikt kg", "Märke / pos"], rp,
         [["Summa per tjocklek", "t mm", "", "m²", "", "", "kg"]] +
         [["", t_, "", round(a_, 2), "", "", round(v, 1)] for t_, (a_, v) in sorted(pertj.items())]),
        ("Per material", ["Material", "Vikt kg"], [[k_, round(v, 1)] for k_, v in sorted(mat.items())], []),
    ]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        Workbook = None
    filer_ut = []
    if Workbook:
        wb = Workbook()
        for i, (titel, kol, data, efter) in enumerate(blad):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = titel
            ws.append(["%s – %s" % (titel.upper(), namn)])
            ws["A1"].font = Font(bold=True, size=12)
            ws.append([GRANSKA])
            ws["A2"].font = Font(italic=True, color="AA0000")
            ws.append([])
            ws.append(kol)
            for c in ws[4]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="DDDDDD")
            for r in data:
                ws.append(r)
            if efter:
                ws.append([])
                for r in efter:
                    ws.append(r)
            for j, k_ in enumerate(kol, 1):
                bredd = max([len(str(k_))] + [len(str(r[j - 1])) for r in data if len(r) >= j])
                ws.column_dimensions[get_column_letter(j)].width = min(max(8, bredd + 2), 60)
            ws.freeze_panes = "A5"
        fil = os.path.join(utmapp, "%s kalkylunderlag.xlsx" % namn)
        wb.save(fil)
        filer_ut.append(fil)
    else:
        for titel, kol, data, efter in blad[:3]:
            fil = os.path.join(utmapp, "%s %s.csv" % (namn, titel.lower()))
            with open(fil, "w", encoding="utf-8-sig") as f:
                f.write(GRANSKA + "\n" + ";".join(kol) + "\n")
                for r in data + ([[]] + efter if efter else []):
                    f.write(";".join(str(x).replace(".", ",") if isinstance(x, float) else str(x) for x in r) + "\n")
            filer_ut.append(fil)
    logg("Kalkylunderlag: %d märken, %d st, %.0f kg, %d kaprader, %d plåtrader -> %s" % (
        len(mangder_), tot[0], tot[2], len(rk), len(rp), ", ".join(os.path.basename(f) for f in filer_ut)))
    return filer_ut


if __name__ == "__fcjobb__":
    import fcapi
    fcapi.kor_jobb(_jobb)
elif __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import fcapi
    if len(sys.argv) > 1 and sys.argv[1] == "samla":
        ap = argparse.ArgumentParser(description="Kalkylunderlag för flera sammanställningar")
        ap.add_argument("filer", nargs="+", help="modell.FCStd eller _mangder.json, valfritt :antal")
        ap.add_argument("--namn", required=True)
        ap.add_argument("--ut", default=None)
        a = ap.parse_args(sys.argv[2:])
        samla(a.filer, a.namn, a.ut)
        sys.exit(0)
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("fil")
    ap.add_argument("--antal", type=int, default=1)
    ap.add_argument("--ut", default=None)
    a = vars(ap.parse_args())
    a["fil"] = os.path.abspath(a["fil"])
    _efterbehandla(fcapi.kor(__file__, a), a)
